"""
Document Parser Utility
Supports parsing DOCX and PDF files to extract text content
"""

import os
from typing import Optional
from docx import Document
import pdfplumber


def parse_docx(file_path: str) -> str:
    """
    Extract text from a DOCX file
    
    Args:
        file_path: Path to the DOCX file
        
    Returns:
        Extracted text content
    """
    try:
        doc = Document(file_path)
        text_content = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text)
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text)
                if row_text:
                    text_content.append(" | ".join(row_text))
        
        return "\n".join(text_content)
    except Exception as e:
        raise Exception(f"Error parsing DOCX file: {str(e)}")


def parse_pdf(file_path: str) -> str:
    """
    Extract text from a PDF file
    
    Args:
        file_path: Path to the PDF file
        
    Returns:
        Extracted text content
    """
    try:
        text_content = []
        
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_content.append(page_text)
        
        return "\n".join(text_content)
    except Exception as e:
        raise Exception(f"Error parsing PDF file: {str(e)}")


def parse_document(file_path: str, file_extension: Optional[str] = None) -> str:
    """
    Auto-detect file type and parse document
    
    Args:
        file_path: Path to the document file
        file_extension: Optional file extension (e.g., '.docx', '.pdf')
        
    Returns:
        Extracted text content
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
    
    # Determine file extension
    if not file_extension:
        _, file_extension = os.path.splitext(file_path)
    
    file_extension = file_extension.lower()
    
    if file_extension == '.docx':
        return parse_docx(file_path)
    elif file_extension == '.pdf':
        return parse_pdf(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_extension}. Supported formats: .docx, .pdf")


def clean_text(text: str) -> str:
    """
    Clean and normalize extracted text
    
    Args:
        text: Raw extracted text
        
    Returns:
        Cleaned text
    """
    # Remove excessive whitespace
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    return '\n'.join(lines)
